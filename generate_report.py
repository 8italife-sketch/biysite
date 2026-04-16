#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Генератор отчета по анализу сайта в формате Word
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

def create_report():
    doc = Document()
    
    # Настройка стилей
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(12)
    
    # Заголовок документа
    title = doc.add_heading('АНАЛИЗ САЙТА И РЕКОМЕНДАЦИИ', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('https://project13514445.tilda.ws/')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True
    
    doc.add_paragraph()
    
    # === ЧАСТЬ 1: ТЕХНИЧЕСКИЙ АУДИТ ===
    h1 = doc.add_heading('1. ТЕХНИЧЕСКИЙ АУДИТ', level=1)
    
    h2 = doc.add_heading('1.1 Критические ошибки (требуют немедленного исправления)', level=2)
    
    critical_issues = [
        ("Карта показывает Нью-Йорк вместо Ставрополя", 
         "Координаты (-74.005941, 40.7127837) неверные. Должны быть: 41.980717, 45.038942 (Ставрополь)."),
        ("Форма заявки не работает", 
         "Атрибут action='' пустой, заявки не отправляются. Требуется настройка обработчика форм."),
        ("Meta-теги шаблонные", 
         "Title: 'Строительная компания', Description: 'Шаблон сайта ремонтной компании'. Требуется уникализация."),
        ("Стандартная фавиконка Tilda", 
         "Отсутствует уникальный брендинг. Необходимо загрузить собственный favicon.ico.")
    ]
    
    for issue, solution in critical_issues:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(issue).bold = True
        doc.add_paragraph(solution, style='List Bullet 2')
    
    doc.add_paragraph()
    
    h2 = doc.add_heading('1.2 Серьёзные проблемы', level=2)
    
    serious_issues = [
        "Отсутствуют реальные кейсы/портфолио",
        "Нет лицензий, сертификатов, информации о компании",
        "Отсутствуют отзывы клиентов",
        "Телефон в шапке не оптимизирован для мобильных устройств"
    ]
    
    for issue in serious_issues:
        doc.add_paragraph(issue, style='List Bullet')
    
    doc.add_paragraph()
    
    h2 = doc.add_heading('1.3 Рекомендации по улучшению', level=2)
    
    doc.add_paragraph('SEO:', style='List Bullet').runs[0].bold = True
    seo_recs = [
        "Добавить robots.txt и sitemap.xml",
        "Прописать alt-тексты для всех изображений",
        "Внедрить микроразметку Schema.org (Organization, LocalBusiness)"
    ]
    for rec in seo_recs:
        doc.add_paragraph(rec, style='List Bullet 2')
    
    doc.add_paragraph()
    doc.add_paragraph('Контент:', style='List Bullet').runs[0].bold = True
    content_recs = [
        "Загрузить реальные фото объектов",
        "Добавить информацию о команде",
        "Указать гарантии и условия работы"
    ]
    for rec in content_recs:
        doc.add_paragraph(rec, style='List Bullet 2')
    
    doc.add_paragraph()
    doc.add_paragraph('UX:', style='List Bullet').runs[0].bold = True
    ux_recs = [
        "Добавить кнопку 'Заказать звонок'",
        "Упростить форму заявки (оставить 2 поля)",
        "Реализовать хлебные крошки"
    ]
    for rec in ux_recs:
        doc.add_paragraph(rec, style='List Bullet 2')
    
    doc.add_paragraph()
    doc.add_paragraph('Скорость:', style='List Bullet').runs[0].bold = True
    speed_recs = [
        "Оптимизировать изображения (формат WebP)",
        "Включить lazy loading для изображений ниже первого экрана"
    ]
    for rec in speed_recs:
        doc.add_paragraph(rec, style='List Bullet 2')
    
    doc.add_paragraph()
    doc.add_paragraph('Доверие:', style='List Bullet').runs[0].bold = True
    trust_recs = [
        "Разместить лицензии СРО",
        "Добавить логотипы партнеров",
        "Указать информацию о страховке объектов"
    ]
    for rec in trust_recs:
        doc.add_paragraph(rec, style='List Bullet 2')
    
    doc.add_page_break()
    
    # === ЧАСТЬ 2: АНАЛИЗ ДИЗАЙНА ===
    h1 = doc.add_heading('2. АНАЛИЗ ДИЗАЙНА И UI/UX', level=1)
    
    h2 = doc.add_heading('2.1 Критические проблемы дизайна', level=2)
    
    doc.add_paragraph('1. Ужасная цветовая схема', style='List Bullet').runs[0].bold = True
    doc.add_paragraph('Проблема: Монохромное использование одного цвета #2986c3 (голубой) для ВСЕХ элементов.', style='List Bullet 2')
    doc.add_paragraph('Решение:', style='List Bullet 2').runs[0].bold = True
    color_table = doc.add_table(rows=5, cols=2)
    color_table.style = 'Table Grid'
    color_data = [
        ['Основной', '#2986c3 (голубой)'],
        ['Акцентный (CTA)', '#FF6B35 (оранжевый)'],
        ['Текст заголовков', '#333333 (темно-серый)'],
        ['Основной текст', '#555555 (серый)'],
        ['Фон секций', '#F8F9FA (светло-серый)']
    ]
    for i, (role, color) in enumerate(color_data):
        color_table.rows[i].cells[0].text = role
        color_table.rows[i].cells[1].text = color
    
    doc.add_paragraph()
    doc.add_paragraph('2. Катастрофическая типографика', style='List Bullet').runs[0].bold = True
    doc.add_paragraph('Проблема: Весь текст голубого цвета #2986c3, низкий контраст, затруднено чтение.', style='List Bullet 2')
    doc.add_paragraph('Решение:', style='List Bullet 2').runs[0].bold = True
    typo_table = doc.add_table(rows=4, cols=2)
    typo_table.style = 'Table Grid'
    typo_data = [
        ['Заголовки', 'Montserrat Bold (700), цвет #333333'],
        ['Подзаголовки', 'Montserrat SemiBold (600), цвет #333333'],
        ['Основной текст', 'Montserrat Regular (400), цвет #555555'],
        ['Ссылки', '#2986c3 с подчеркиванием при hover']
    ]
    for i, (element, spec) in enumerate(typo_data):
        typo_table.rows[i].cells[0].text = element
        typo_table.rows[i].cells[1].text = spec
    
    doc.add_paragraph()
    doc.add_paragraph('3. Отсутствие воздушности (whitespace)', style='List Bullet').runs[0].bold = True
    doc.add_paragraph('Решение: Увеличить padding секций до 120px на десктопе, добавить светло-серые фоны для чередования секций, увеличить межстрочный интервал до 1.6-1.8.', style='List Bullet 2')
    
    doc.add_paragraph()
    
    h2 = doc.add_heading('2.2 Серьёзные проблемы', level=2)
    
    issues_serious = [
        ("Слабый первый экран (Hero Section)", 
         ["Заголовок 'Строительная компания' слишком шаблонный", 
          "Отсутствует УТП (уникальное торговое предложение)",
          "Рекомендация: 'Строим промышленные объекты под ключ в Ставрополе и ЮФО'"]),
        ("Непрофессиональные изображения", 
         ["Отсутствуют реальные фото объектов", 
          "Нет фото команды и техники",
          "Рекомендация: Загрузить минимум 10-15 реальных фото"]),
        ("Ужасная форма заявки", 
         ["action='' (форма не работает!)", 
          "3 обязательных поля (слишком много)",
          "Placeholder '+1(000)000-0000' (американский формат для России!)",
          "Рекомендация: Оставить 2 поля (Имя + Телефон), маска +7 (___) ___-__-__"]),
        ("Карта с Нью-Йорком", 
         ["Текущие координаты: -74.005941, 40.7127837 (Нью-Йорк!)", 
          "Должно быть: 41.980717, 45.038942 (Ставрополь)"])
    ]
    
    for title, points in issues_serious:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(title).bold = True
        for point in points:
            doc.add_paragraph(point, style='List Bullet 2')
    
    doc.add_page_break()
    
    h2 = doc.add_heading('2.3 Проблемы UX/UI', level=2)
    
    ux_issues = [
        ("Навигация", 
         ["Закрепить шапку (position: fixed)", 
          "Добавить плавный скролл к якорям",
          "Реализовать бургер-меню для мобильных"]),
        ("Блок 'Этапы работы'", 
         ["Перестроить в 2 ряда (3 + 2) или вертикальный timeline", 
          "Добавить иконки для каждого этапа",
          "Выделить сроки крупным шрифтом"]),
        ("FAQ (Вопрос-ответ)", 
         ["Добавить вопросы: 'Какие гарантии вы предоставляете?'", 
          "'Можно ли платить поэтапно?'",
          "'Что входит в смету?'"]),
        ("Контакты", 
         ["Исправить координаты карты", 
          "Добавить график работы: 'Пн-Пт 9:00-18:00'",
          "Кнопки WhatsApp, Telegram",
          "Полный адрес: 'офис 205'"])
    ]
    
    for title, points in ux_issues:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(title).bold = True
        for point in points:
            doc.add_paragraph(point, style='List Bullet 2')
    
    doc.add_paragraph()
    
    h2 = doc.add_heading('2.4 Технические спецификации', level=2)
    
    doc.add_paragraph('Цветовая палитра (CSS переменные):', style='List Bullet').runs[0].bold = True
    css_code = """:root {
  --primary: #2986c3;        /* Голубой - бренд */
  --accent: #FF6B35;         /* Оранжевый - CTA */
  --dark: #333333;           /* Текст заголовков */
  --gray: #555555;           /* Основной текст */
  --light-gray: #F8F9FA;     /* Фон секций */
  --white: #FFFFFF;          /* Фон */
  --success: #28a745;        /* Для отзывов, гарантий */
}"""
    doc.add_paragraph(css_code).style = 'No Spacing'
    
    doc.add_paragraph()
    doc.add_paragraph('Типографика:', style='List Bullet').runs[0].bold = True
    typo_code = """h1 { font-size: 48px; font-weight: 700; color: #333; line-height: 1.2; }
h2 { font-size: 36px; font-weight: 600; color: #333; line-height: 1.3; }
h3 { font-size: 24px; font-weight: 600; color: #333; line-height: 1.4; }
p  { font-size: 16px; font-weight: 400; color: #555; line-height: 1.7; }"""
    doc.add_paragraph(typo_code).style = 'No Spacing'
    
    doc.add_paragraph()
    doc.add_paragraph('Кнопки CTA:', style='List Bullet').runs[0].bold = True
    btn_code = """.btn-primary {
  background: #FF6B35;
  color: #FFFFFF;
  padding: 18px 36px;
  font-size: 16px;
  font-weight: 600;
  border-radius: 5px;
  transition: all 0.3s;
}
.btn-primary:hover {
  background: #E55A2B;
  transform: translateY(-2px);
  box-shadow: 0 4px 12px rgba(255,107,53,0.4);
}"""
    doc.add_paragraph(btn_code).style = 'No Spacing'
    
    doc.add_page_break()
    
    # === ЧАСТЬ 3: ПРИОРИТЕТЫ И ПЛАН ===
    h1 = doc.add_heading('3. ПЛАНЫ И ПРИОРИТЕТЫ', level=1)
    
    h2 = doc.add_heading('3.1 Быстрые победы (сделать за 1 день)', level=2)
    
    quick_wins = [
        ("Исправить карту", "5 мин"),
        ("Поменять цвет текста с голубого на темно-серый", "30 мин"),
        ("Настроить форму (action, 2 поля, маска телефона)", "30 мин"),
        ("Переписать заголовок первого экрана (УТП)", "20 мин"),
        ("Добавить favicon", "10 мин"),
        ("Поменять текст кнопки CTA", "5 мин")
    ]
    
    table = doc.add_table(rows=len(quick_wins)+1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Задача'
    hdr_cells[1].text = 'Время'
    hdr_cells[2].text = 'Приоритет'
    
    for i, (task, time) in enumerate(quick_wins):
        row_cells = table.rows[i+1].cells
        row_cells[0].text = task
        row_cells[1].text = time
        row_cells[2].text = 'P0'
    
    doc.add_paragraph()
    
    h2 = doc.add_heading('3.2 Приоритеты дизайн-улучшений', level=2)
    
    priorities = [
        ("P0", "Исправить карту", "5 мин", "Высокое"),
        ("P0", "Поменять цвет текста", "30 мин", "Критичное"),
        ("P0", "Настроить форму", "30 мин", "Критичное"),
        ("P0", "Переписать УТП", "20 мин", "Высокое"),
        ("P1", "Добавить реальные фото", "2 часа", "Высокое"),
        ("P1", "Улучшить цветовую схему", "1 час", "Среднее"),
        ("P2", "Переработать первый экран", "2 часа", "Высокое"),
        ("P2", "Добавить отзывы", "1 час", "Среднее"),
        ("P3", "Анимации и микро-UX", "3 часа", "Низкое")
    ]
    
    table = doc.add_table(rows=len(priorities)+1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Приоритет'
    hdr_cells[1].text = 'Задача'
    hdr_cells[2].text = 'Время'
    hdr_cells[3].text = 'Влияние'
    
    for i, (prio, task, time, impact) in enumerate(priorities):
        row_cells = table.rows[i+1].cells
        row_cells[0].text = prio
        row_cells[1].text = task
        row_cells[2].text = time
        row_cells[3].text = impact
    
    doc.add_paragraph()
    
    h2 = doc.add_heading('3.3 Дополнительные идеи', level=2)
    
    ideas = [
        "Квиз: 'Рассчитайте стоимость строительства за 5 вопросов'",
        "Калькулятор: Примерный расчет стоимости м²",
        "Видео: Ролик о компании (1-2 мин)",
        "Сертификаты: Скан-копии лицензий СРО",
        "Команда: Фото ключевых сотрудников с должностями"
    ]
    
    for idea in ideas:
        doc.add_paragraph(idea, style='List Bullet')
    
    doc.add_paragraph()
    
    # Заключение
    h2 = doc.add_heading('4. ЗАКЛЮЧЕНИЕ', level=2)
    
    conclusion = """Сайт выглядит как шаблон 2015 года. Требуется полная переработка визуального стиля, контента и UX для соответствия современным стандартам строительной индустрии 2026 года.

Все выявленные критические ошибки могут быть исправлены в течение 1 рабочего дня силами редактора Tilda без привлечения разработчиков.

Реализация всех рекомендаций позволит повысить конверсию сайта в 3-5 раз и улучшить доверие потенциальных клиентов."""
    
    doc.add_paragraph(conclusion)
    
    # Сохранение файла
    file_path = '/workspace/website_analysis_report.docx'
    doc.save(file_path)
    print(f"Отчет успешно создан: {file_path}")
    return file_path

if __name__ == '__main__':
    create_report()
